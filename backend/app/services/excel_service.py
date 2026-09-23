# -*- coding: utf-8 -*-
"""Excel 报表导出服务（openpyxl，内存生成，风格统一）"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F6FB2")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)

# Excel 公式注入防护：openpyxl 会把以 = 开头的字符串当公式写入，
# 导出内容含用户可控文本（审核备注/车牌等）时可能触发 DDE/公式执行
_RISKY_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_cell(value):
    if isinstance(value, str) and value.startswith(_RISKY_PREFIXES):
        return "'" + value
    return value


def build_xlsx(title: str, headers: list[str], rows: list[list]) -> bytes:
    """生成统一风格的 xlsx 字节流：标题行 + 表头 + 数据，自动列宽"""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else "Sheet1"

    # 标题行（合并单元格）
    ws.cell(row=1, column=1, value=f"{title}（导出时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}）")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    ws.cell(row=1, column=1).font = Font(bold=True, size=13, color="1F6FB2")
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # 表头
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 数据行
    for r, row in enumerate(rows, start=3):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=_safe_cell(value))

    # 自动列宽（按前 200 行采样，中文按 2 列宽估算）
    for col in range(1, len(headers) + 1):
        width = len(str(headers[col - 1])) + 4
        for row in rows[:200]:
            value = "" if col > len(row) else row[col - 1]
            width = max(width, sum(2 if ord(ch) > 127 else 1 for ch in str(value)) + 2)
        ws.column_dimensions[get_column_letter(col)].width = min(42, width)

    # 冻结表头
    ws.freeze_panes = "A3"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
