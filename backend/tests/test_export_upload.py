# -*- coding: utf-8 -*-
"""优化迭代功能测试：Excel 报表导出、通用图片上传、流量数据保留清理"""
import io
import os
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app.core.config import settings
from app.models import TrafficFlow
from app.tasks.scheduler import cleanup_flow_data

UPLOAD_DIR = settings.UPLOAD_DIR


def _tiny_png() -> bytes:
    """用 PIL 生成 1x1 PNG（openpyxl/PIL 均在依赖内）"""
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (1, 1), (200, 60, 60)).save(buf, format="PNG")
    return buf.getvalue()


class TestExcelExport:
    def test_export_violations(self, client, admin_headers):
        resp = client.get("/api/v1/export/violations.xlsx", headers=admin_headers)
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.cell(row=2, column=2).value == "车牌号"

    def test_export_parking_records(self, client, admin_headers):
        resp = client.get("/api/v1/export/parking-records.xlsx", headers=admin_headers)
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        assert wb.active.cell(row=2, column=3).value == "车牌号"

    def test_export_traffic_report(self, client, admin_headers):
        resp = client.get("/api/v1/export/traffic-report.xlsx", headers=admin_headers)
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        headers = [wb.active.cell(row=2, column=c).value for c in range(1, 7)]
        assert headers[0] == "日期" and headers[-1] == "拥堵占比"

    def test_export_forbidden_for_user(self, client, user_headers):
        assert client.get("/api/v1/export/violations.xlsx", headers=user_headers).status_code == 403
        assert client.get("/api/v1/export/traffic-report.xlsx", headers=user_headers).status_code == 403


class TestImageUpload:
    def test_upload_image_ok(self, client, officer_headers):
        resp = client.post(
            "/api/v1/upload/image",
            headers=officer_headers,
            files={"file": ("evidence.png", _tiny_png(), "image/png")},
        )
        body = resp.json()
        assert body["code"] == 0
        url = body["data"]["url"]
        assert url.startswith("/static/uploads/") and url.endswith(".png")
        # 文件确实落盘（uploads 相对 backend 运行目录）
        local = os.path.join(UPLOAD_DIR, os.path.basename(url))
        assert os.path.exists(local)
        os.remove(local)  # 清理测试产物

    def test_upload_bad_suffix(self, client, admin_headers):
        resp = client.post(
            "/api/v1/upload/image",
            headers=admin_headers,
            files={"file": ("evil.exe", b"MZ...", "application/octet-stream")},
        )
        assert resp.json()["code"] == 90002

    def test_upload_forbidden_for_user(self, client, user_headers):
        resp = client.post(
            "/api/v1/upload/image",
            headers=user_headers,
            files={"file": ("a.png", _tiny_png(), "image/png")},
        )
        assert resp.status_code == 403


class TestFlowRetention:
    def test_cleanup_deletes_only_expired(self, client, db_session, monkeypatch):
        """保留策略：只清理超过保留期的时序数据，近期数据不受影响"""
        monkeypatch.setattr(settings, "FLOW_RETENTION_DAYS", 30)
        now = datetime.now()
        old = TrafficFlow(
            road_section_id=1,
            recorded_at=now - timedelta(days=40),
            flow=100,
            speed=40,
            saturation=0.2,
            congestion_level=0,
        )
        fresh = TrafficFlow(
            road_section_id=1,
            recorded_at=now,
            flow=200,
            speed=45,
            saturation=0.3,
            congestion_level=0,
        )
        db_session.add_all([old, fresh])
        db_session.commit()
        old_id, fresh_id = old.id, fresh.id

        deleted = cleanup_flow_data(db=db_session)
        assert deleted >= 1
        db_session.expire_all()  # 清理走 synchronize_session=False，先过期身份映射
        assert db_session.get(TrafficFlow, old_id) is None
        assert db_session.get(TrafficFlow, fresh_id) is not None

    def test_cleanup_disabled_when_zero(self, client, db_session, monkeypatch):
        monkeypatch.setattr(settings, "FLOW_RETENTION_DAYS", 0)
        assert cleanup_flow_data(db=db_session) == 0


class TestUploadContentValidation:
    def test_fake_image_rejected(self, client, officer_headers):
        """改后缀伪装：jpg 扩展名 + HTML 内容必须被 magic bytes 校验拦截"""
        resp = client.post(
            "/api/v1/upload/image",
            headers=officer_headers,
            files={"file": ("fake.jpg", b"<html>definitely not an image</html>", "image/jpeg")},
        )
        assert resp.json()["code"] == 90002


class TestExcelFormulaSanitize:
    def test_formula_injection_neutralized(self):
        """导出内容含用户可控文本时，= 开头的字符串不得被 openpyxl 当公式写入"""
        from app.services.excel_service import build_xlsx

        content = build_xlsx("测试", ["备注"], [["=cmd|'/c calc'!A1"], ["正常文本"]])
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        assert str(ws.cell(row=3, column=1).value).startswith("'")
        assert ws.cell(row=4, column=1).value == "正常文本"
